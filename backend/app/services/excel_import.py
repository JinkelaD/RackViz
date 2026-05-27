from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from ..models.rack import Rack
from ..models.device import Device
from ..models.device_model import DeviceModel

def import_devices_excel(file_bytes: bytes, db: Session) -> dict:
    wb = load_workbook(filename=BytesIO(file_bytes))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    imported = 0
    errors = []

    for row_idx, row in enumerate(rows):
        if row_idx < 2:
            continue
        if not row or not row[0]:
            continue
        try:
            name = str(row[0]).strip()
            model_name = str(row[1]).strip() if len(row) > 1 and row[1] else ""
            device_type = str(row[2]).strip() if len(row) > 2 and row[2] else "server"
            location = str(row[3]).strip() if len(row) > 3 and row[3] else ""
            manufacturer = str(row[4]).strip() if len(row) > 4 and row[4] else ""
            function = str(row[5]).strip() if len(row) > 5 and row[5] else ""

            rack_name = ""
            start_u = None
            end_u = None
            if location:
                parts = location.replace("，", ",").split(",")
                if len(parts) >= 1:
                    rack_name = parts[0].strip()
                if len(parts) >= 2:
                    u_part = parts[1].strip().replace("U", "").replace("u", "")
                    if "-" in u_part:
                        u_parts = u_part.split("-")
                        try:
                            start_u = int(u_parts[0])
                            end_u = int(u_parts[1])
                        except ValueError:
                            pass

            rack = None
            if rack_name:
                rack = db.query(Rack).filter(Rack.name == rack_name).first()
                if not rack:
                    rack = Rack(name=rack_name, height_u=42)
                    db.add(rack)
                    db.flush()

            model = None
            if model_name:
                model = db.query(DeviceModel).filter(DeviceModel.name == model_name).first()
                if not model:
                    model = DeviceModel(
                        name=model_name, type=device_type,
                        manufacturer=manufacturer, height_u=1,
                    )
                    db.add(model)
                    db.flush()

            device = Device(
                name=name,
                device_model_id=model.id if model else None,
                rack_id=rack.id if rack else None,
                start_u=start_u,
                end_u=end_u,
                function=function,
            )
            db.add(device)
            imported += 1
        except Exception as e:
            errors.append(f"Row {row_idx + 1}: {str(e)}")

    db.commit()
    return {"imported": imported, "errors": errors}