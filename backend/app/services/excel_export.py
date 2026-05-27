from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from ..models.rack import Rack
from ..models.device import Device

U_HEIGHT = 22

type_colors = {
    "server": "5A8FD4",
    "switch": "3DC9B0",
    "router": "D4A85A",
    "storage": "8B5AD4",
    "pdu": "D45A5A",
    "patch": "5C6170",
}

def export_racks_excel(db: Session) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "机柜部署图"

    racks = db.query(Rack).order_by(Rack.row, Rack.col).all()
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    col_offset = 0
    for rack in racks:
        if col_offset == 0:
            ws.column_dimensions["A"].width = 6
            for u in range(rack.height_u, 0, -1):
                cell = ws.cell(row=rack.height_u - u + 1, column=1, value=f"{u}U")
                cell.font = Font(name="Consolas", size=8, color="888888")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                ws.row_dimensions[rack.height_u - u + 1].height = U_HEIGHT

        header_col = col_offset + 2
        ws.merge_cells(start_row=1, start_column=header_col, end_row=1, end_column=header_col)
        header_cell = ws.cell(row=1, column=header_col, value=rack.name)
        header_cell.font = Font(name="Microsoft YaHei", size=10, bold=True)
        header_cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(header_col)].width = 24

        sub_cell = ws.cell(row=2, column=header_col, value=f"{rack.height_u}U")
        sub_cell.font = Font(name="Consolas", size=8, color="888888")
        sub_cell.alignment = Alignment(horizontal="center")

        devices = db.query(Device).filter(Device.rack_id == rack.id).all()
        for u in range(rack.height_u, 0, -1):
            row_num = rack.height_u - u + 1
            cell = ws.cell(row=row_num, column=header_col)
            cell.border = thin_border

        for device in devices:
            if device.start_u is None or device.end_u is None:
                continue
            model = device.device_model
            color = type_colors.get(model.type if model else "", "5A8FD4")
            fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            font = Font(name="Microsoft YaHei", size=8, color="FFFFFF", bold=True)

            for u in range(device.start_u, device.end_u + 1):
                row_num = rack.height_u - u + 1
                cell = ws.cell(row=row_num, column=header_col, value=device.name)
                cell.fill = fill
                cell.font = font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

        col_offset += 1

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output